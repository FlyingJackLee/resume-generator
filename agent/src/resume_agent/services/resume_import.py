from __future__ import annotations

import io
import logging
import re
import time
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import yaml
from pydantic import BaseModel, ConfigDict, Field

from resume_agent.errors import ResumeAgentError

Language = Literal["zh", "en"]
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}
logger = logging.getLogger(__name__)

# A few stray characters from a PDF's metadata or hidden layer are not a resume.
# Do this check before an extracted string is ever passed to the LLM.
MIN_RESUME_TEXT_CHARACTERS = 50


class ParsedResume(BaseModel):
    model_config = ConfigDict(extra="forbid")
    resume: dict[str, Any]
    source_languages: list[Language] = Field(min_length=1, max_length=2)
    warnings: list[str] = Field(default_factory=list, max_length=20)


IMPORT_SYSTEM_PROMPT = """Convert the supplied resume into this application's baseline resume data, using EXACTLY this JSON shape for the "resume" field. Every text value is a bilingual object {"zh": "...", "en": "..."}.

{
  "meta": {
    "name": {"zh": "", "en": ""},
    "position": {"zh": "", "en": ""},
    "contacts": {"mobile": "", "email": "", "homepage": "", "github": "", "wechat": ""}
  },
  "sections": [
    {"type": "paragraph", "title": {"zh": "", "en": ""}, "body": {"zh": "", "en": ""}},
    {"type": "skills", "title": {"zh": "", "en": ""}, "rows": [
      {"label": {"zh": "", "en": ""}, "items": {"zh": "", "en": ""}}
    ]},
    {"type": "education", "title": {"zh": "", "en": ""}, "entries": [
      {"title": {"zh": "", "en": ""}, "org": {"zh": "", "en": ""}, "location": {"zh": "", "en": ""}, "date": {"zh": "", "en": ""}}
    ]},
    {"type": "entries", "title": {"zh": "", "en": ""}, "entries": [
      {"title": {"zh": "", "en": ""}, "org": {"zh": "", "en": ""}, "location": {"zh": "", "en": ""}, "date": {"zh": "", "en": ""},
       "summary": {"zh": "", "en": ""}, "items": [{"zh": "", "en": ""}]}
    ]}
  ]
}

Rules:
- Every section MUST use one of these four "type" values and carry exactly its listed keys. Never invent alternative keys
  (e.g. a top-level "items" list on a section, "name"/"description" on a skill row, or a plain-string "title").
- Use "entries" sections for work experience, projects, and certificates/honors alike; put bullet points in that entry's
  "items" (one bilingual object per bullet), not "responsibilities" or a plain string.
- Use "education" only for the schooling section; put degree/major in that entry's "title".
- Extract only visible source facts. Never invent or translate content.
- If source is Chinese only, use zh and leave en empty; if English only, use en and leave zh empty.
- Put anything uncertain or that does not fit this schema into "warnings" instead of dropping it silently."""


def source_language_hint(text: str) -> list[Language] | None:
    zh, latin = bool(re.search(r"[\u3400-\u9fff]", text)), bool(re.search(r"[A-Za-z]", text))
    # A Chinese resume commonly contains English product names, while a bilingual
    # resume naturally contains both scripts. Only override the model for unambiguous text.
    return None if zh and latin else ["zh"] if zh else ["en"]


def has_meaningful_resume_text(text: str) -> bool:
    """Reject empty, garbled, and metadata-only extraction results."""
    visible = "".join(character for character in text if character.isprintable())
    if not visible or visible.count("\ufffd") > 2:
        return False
    meaningful = re.findall(r"[\u3400-\u9fffA-Za-z0-9]", visible)
    return len(meaningful) >= MIN_RESUME_TEXT_CHARACTERS


def require_meaningful_resume_text(text: str, *, failure_message: str) -> str:
    cleaned = text.strip()
    if not has_meaningful_resume_text(cleaned):
        raise ResumeAgentError(failure_message)
    return cleaned


def _bi(value: Any) -> dict[str, str]:
    return {"zh": str(value.get("zh") or ""), "en": str(value.get("en") or "")} if isinstance(value, dict) else {"zh": "", "en": ""}


def _erase_missing(value: Any, missing: set[Language]) -> None:
    if isinstance(value, dict):
        if "zh" in value or "en" in value:
            for language in missing:
                value[language] = ""
        for child in value.values():
            _erase_missing(child, missing)
    elif isinstance(value, list):
        for child in value:
            _erase_missing(child, missing)


def normalize_resume(candidate: dict[str, Any], languages: list[Language]) -> dict[str, Any]:
    if not isinstance(candidate, dict):
        raise ResumeAgentError("解析结果不是简历对象")
    meta = candidate.get("meta") if isinstance(candidate.get("meta"), dict) else {}
    contacts = meta.get("contacts") if isinstance(meta.get("contacts"), dict) else {}
    result: dict[str, Any] = {
        "meta": {
            "name": _bi(meta.get("name")), "name_split_zh": int(meta.get("name_split_zh") or 1),
            "position": _bi(meta.get("position")), "color": str(meta.get("color") or "#2563EB"),
            "contacts": {key: str(contacts.get(key) or "") for key in ("mobile", "email", "homepage", "github", "wechat")},
            "photo": None, "footer_label": _bi(meta.get("footer_label")),
            "footer_more": {"text": _bi((meta.get("footer_more") or {}).get("text") if isinstance(meta.get("footer_more"), dict) else None), "url": str((meta.get("footer_more") or {}).get("url") or "") if isinstance(meta.get("footer_more"), dict) else ""},
        },
        "section_highlight_latin": int(candidate.get("section_highlight_latin") or 0),
        "section_highlight_cjk": int(candidate.get("section_highlight_cjk") or 0),
        "sections": candidate.get("sections") if isinstance(candidate.get("sections"), list) else [],
    }
    for section in result["sections"]:
        if not isinstance(section, dict):
            raise ResumeAgentError("解析结果包含无效简历模块")
        section.pop("id", None); section.pop("facts", None); section["title"] = _bi(section.get("title"))
        if section.get("type") == "paragraph": section["body"] = _bi(section.get("body"))
        for row in section.get("rows", []):
            if isinstance(row, dict): row.pop("id", None); row["label"] = _bi(row.get("label")); row["items"] = _bi(row.get("items"))
        for entry in section.get("entries", []):
            if isinstance(entry, dict):
                entry.pop("id", None); entry.pop("facts", None)
                for key in ("title", "org", "location", "date", "summary"):
                    if key in entry: entry[key] = _bi(entry.get(key))
                for key in ("items", "responsibilities"):
                    value = entry.get(key)
                    if value is not None:
                        values = value if isinstance(value, list) else [value]
                        entry[key] = [_bi(item) for item in values if isinstance(item, dict)]
    _erase_missing(result, {"zh", "en"} - set(languages))
    return result


@lru_cache(maxsize=1)
def _ocr_engine() -> Any:
    try:
        from rapidocr import RapidOCR

        return RapidOCR()
    except Exception as exc:
        raise ResumeAgentError("本地 RapidOCR 初始化失败；请执行 uv sync 后重试") from exc


def _ocr_image(data: bytes) -> str:
    try:
        result = _ocr_engine()(data)
    except ResumeAgentError:
        raise
    except Exception as exc:
        raise ResumeAgentError("本地 OCR 解析失败，请上传更清晰的图片") from exc
    return "\n".join(result.txts or ()).strip()


def extract_local_ocr(
    filename: str,
    data: bytes,
    *,
    max_pages: int,
) -> str:
    """OCR scanned PDFs and images locally; no original file leaves this machine."""
    suffix = Path(filename).suffix.lower()
    if suffix in IMAGE_TYPES:
        text = _ocr_image(data)
        return require_meaningful_resume_text(text, failure_message="本地 OCR 未识别到足够的简历文字，请上传更清晰的图片")
    if suffix != ".pdf":
        raise ResumeAgentError("本地 OCR 仅用于 PDF 或图片")
    try:
        import fitz

        document = fitz.open(stream=data, filetype="pdf")
        pages = [page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png") for page in document[:max_pages]]
        document.close()
    except Exception as exc:
        raise ResumeAgentError("扫描 PDF 无法转换为本地 OCR 图片") from exc
    if not pages:
        raise ResumeAgentError("PDF 不含可识别页面")
    text = "\n\n".join(
        _ocr_image(page)
        for page in pages
    ).strip()
    return require_meaningful_resume_text(text, failure_message="本地 OCR 未识别到足够的简历文字，请上传更清晰的扫描件")


def _mineru_request(url: str, *, method: str, data: bytes | None = None, token: str = "") -> Any:
    headers = {"Accept": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if data is not None:
        headers["Content-Type"] = "application/json"
    try:
        with urlopen(Request(url, data=data, headers=headers, method=method), timeout=30) as response:
            return response.read()
    except (HTTPError, URLError) as exc:
        raise ResumeAgentError("MinerU API 请求失败，请检查网络、Token 和服务状态") from exc


def extract_mineru_api(
    filename: str,
    data: bytes,
    *,
    base_url: str,
    token: str,
    model_version: str,
    timeout_seconds: int,
) -> str:
    """Use MinerU's signed-upload API and return full.md from its result archive."""
    if not token:
        raise ResumeAgentError("已选择 MinerU API，但未设置 RESUME_AGENT_MINERU_API_TOKEN")
    base_url = base_url.rstrip("/")
    payload = {
        "files": [{"name": filename, "data_id": "resume-import", "is_ocr": True}],
        "model_version": model_version,
        "language": "ch",
        "enable_table": True,
        "enable_formula": False,
    }
    import json

    response = json.loads(_mineru_request(
        f"{base_url}/file-urls/batch", method="POST", data=json.dumps(payload).encode(), token=token
    ))
    if response.get("code") != 0:
        raise ResumeAgentError(f"MinerU API 创建任务失败：{response.get('msg', '未知错误')}")
    details = response.get("data") or {}
    upload_urls = details.get("file_urls") or []
    batch_id = details.get("batch_id")
    if not batch_id or len(upload_urls) != 1:
        raise ResumeAgentError("MinerU API 未返回有效上传任务")
    try:
        with urlopen(Request(upload_urls[0], data=data, method="PUT"), timeout=60) as upload:
            if upload.status not in {200, 201, 204}:
                raise ResumeAgentError("MinerU API 上传简历失败")
    except (HTTPError, URLError) as exc:
        raise ResumeAgentError("MinerU API 上传简历失败") from exc

    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        result = json.loads(_mineru_request(
            f"{base_url}/extract-results/batch/{batch_id}", method="GET", token=token
        ))
        items = (result.get("data") or {}).get("extract_result") or []
        item = items[0] if items else {}
        state = item.get("state")
        if state == "done" and item.get("full_zip_url"):
            archive = _mineru_request(item["full_zip_url"], method="GET")
            try:
                with zipfile.ZipFile(io.BytesIO(archive)) as bundle:
                    markdown_name = next(name for name in bundle.namelist() if Path(name).name == "full.md")
                    text = bundle.read(markdown_name).decode("utf-8").strip()
            except (OSError, StopIteration, UnicodeDecodeError, zipfile.BadZipFile) as exc:
                raise ResumeAgentError("MinerU API 返回结果中缺少可读取的 Markdown") from exc
            return require_meaningful_resume_text(text, failure_message="MinerU API 未解析出足够的简历文字")
        if state == "failed":
            raise ResumeAgentError(f"MinerU API 解析失败：{item.get('err_msg', '未知错误')}")
        time.sleep(2)
    raise ResumeAgentError("MinerU API 解析超时，请稍后重试或改用 RapidOCR")


def extract_upload(
    filename: str,
    data: bytes,
    *,
    ocr_max_pages: int = 5,
    scan_parser: str = "rapidocr",
    mineru_api_base_url: str = "https://mineru.net/api/v4",
    mineru_api_token: str = "",
    mineru_model_version: str = "vlm",
    mineru_timeout_seconds: int = 180,
) -> tuple[str | None, dict[str, Any] | None, str]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".yaml", ".yml"}:
        try: value = yaml.safe_load(data.decode("utf-8"))
        except (UnicodeDecodeError, yaml.YAMLError) as exc: raise ResumeAgentError("YAML 简历无法读取，请使用 UTF-8 编码和有效 YAML") from exc
        if not isinstance(value, dict): raise ResumeAgentError("YAML 简历根节点必须是对象")
        return None, value, "YAML 结构读取"
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages).strip()
        except Exception as exc: raise ResumeAgentError("PDF 无法读取或已加密") from exc
        if has_meaningful_resume_text(text):
            logger.info("Resume import extracted embedded PDF text chars=%d", len(text))
            return text.strip(), None, "PDF 内嵌文本提取"
        logger.info("Resume import PDF has no usable embedded text; starting %s preprocessing", scan_parser)
        text = extract_mineru_api(filename, data, base_url=mineru_api_base_url, token=mineru_api_token, model_version=mineru_model_version, timeout_seconds=mineru_timeout_seconds) if scan_parser == "mineru_api" else extract_local_ocr(filename, data, max_pages=ocr_max_pages)
        return text, None, "MinerU OCR 预处理" if scan_parser == "mineru_api" else "本地 OCR 预处理"
    if suffix == ".docx":
        try:
            from docx import Document
            document = Document(io.BytesIO(data))
        except Exception as exc: raise ResumeAgentError("Word 文件无法读取，请使用 .docx 格式") from exc
        text = "\n".join(part for part in [*(p.text for p in document.paragraphs), *(" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows)] if part.strip())
        return require_meaningful_resume_text(
            text,
            failure_message="该 Word 文件没有可提取的简历文字（可能仅包含图片）。请先另存为 PDF 后重新上传，或改用 YAML。",
        ), None, "Word 文本提取"
    if suffix in IMAGE_TYPES:
        text = extract_mineru_api(filename, data, base_url=mineru_api_base_url, token=mineru_api_token, model_version=mineru_model_version, timeout_seconds=mineru_timeout_seconds) if scan_parser == "mineru_api" else extract_local_ocr(filename, data, max_pages=ocr_max_pages)
        return text, None, "MinerU OCR 预处理" if scan_parser == "mineru_api" else "本地 OCR 预处理"
    raise ResumeAgentError("仅支持 PDF、DOCX、PNG、JPG、WEBP 或 YAML 简历")
